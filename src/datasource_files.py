"""CSV/Excel data-source connector (pandas). The file is uploaded and stored
under DATA_DIR/qb_sources/<id>/; its path lives in DataSource.options JSON.
pandas/openpyxl are lazy-imported so the core app runs without them.

v2 (23-05, QB-17): `fetch_table` streams instead of materializing the whole
file into one in-memory list. CSV: a bounded first sample (`CSV_SAMPLE_ROWS`
via `pd.read_csv(nrows=...)`) infers columns/types, then the SyncTable's
`rows` generator streams the file in `CSV_CHUNK_ROWS`-row chunks
(`pd.read_csv(chunksize=...)`), pinning `dtype=` to the RESOLVED column type
(`sync_table.columns[i].type_`, mutated in place by Plan-03's
`resolve_effective_types()` between `fetch_table()` and
`materialize_synctable()`) read LAZILY at generator-consumption time -- so
every chunk parses consistently instead of pandas independently (and
inconsistently) inferring a dtype per chunk (RESEARCH Pitfall 4). Excel has
no pandas chunksize; it streams via `openpyxl.load_workbook(read_only=True)`
+ `ws.iter_rows(values_only=True)`, inferring columns/types from a bounded
first sample (`EXCEL_SAMPLE_ROWS`, `itertools.islice`/`chain` -- mirroring
the REST connector's own bounded-sample pattern, 23-04). Both branches track
approximate bytes streamed against `MAX_SYNC_BYTES` (imported from
`datasource_connectors`, the same constant REST's byte budget uses) and set
`SyncTable.truncated = True` (MD-01 reuse) once the cap fires, discarding
the chunk/row that would push over budget rather than partially yielding
it. `FileConnector.__init__`'s path confinement is completely untouched."""
import itertools
import json
import logging
import os
from datetime import date, datetime as _datetime
from pathlib import PurePath

from sqlalchemy.types import Integer, Float, Boolean, DateTime, String

from core.constants import DATA_DIR
from src.datasource_connectors import (
    DataSourceConnector, SyncColumn, SyncTable, ConnectorError, MissingDriverError,
    MAX_SYNC_BYTES,
)

logger = logging.getLogger(__name__)

# QB-17/D-17: bounded first-sample sizes used to infer columns/types before
# streaming the remainder -- a generator can't be scanned in full up-front
# without defeating the whole point of streaming.
CSV_SAMPLE_ROWS = 500
EXCEL_SAMPLE_ROWS = 500
# pd.read_csv(chunksize=...) -- mirrors materialize_synctable's own _BATCH
# scale (1000), so one CSV chunk lines up with roughly one materialize flush.
CSV_CHUNK_ROWS = 1000


def _pandas():
    try:
        import pandas as pd
        return pd
    except ImportError as e:
        raise MissingDriverError("CSV/Excel sources need pandas — pip install pandas openpyxl") from e


def _openpyxl():
    """Lazy-import guard for Excel streaming, mirroring `_pandas()`'s idiom:
    surfaces a MissingDriverError with an install hint instead of a raw
    ImportError, so a source referencing an Excel sheet without openpyxl
    installed fails the same clean way `_pandas()` does."""
    try:
        import openpyxl
        return openpyxl
    except ImportError as e:
        raise MissingDriverError("Excel sources need openpyxl — pip install openpyxl") from e


def _dtype_to_sa(dtype):
    s = str(dtype)
    if s.startswith("int"):
        return Integer()
    if s.startswith("float"):
        return Float()
    if s == "bool":
        return Boolean()
    if s.startswith("datetime"):
        return DateTime()
    return String()


def _sa_type_to_pandas_dtype(type_) -> str:
    """Map a RESOLVED SQLAlchemy column type back to a pandas dtype name for
    `pd.read_csv(dtype=...)`, so every chunk parses that column consistently
    instead of pandas independently inferring (and potentially drifting) a
    dtype per chunk (QB-17/D-15, RESEARCH Pitfall 4). Nullable "Int64"/
    "boolean" dtypes tolerate a missing/NaN cell without forcing an unwanted
    float promotion that would itself look like drift. DateTime/String both
    fall through to "object" (kept as raw text) -- `dtype=` does not reliably
    parse date strings the way `parse_dates=` does, and QB-10's contract is
    type STABILITY across chunks, not additional value-level parsing."""
    if isinstance(type_, Integer):
        return "Int64"
    if isinstance(type_, Float):
        return "float64"
    if isinstance(type_, Boolean):
        return "boolean"
    return "object"


def _infer_excel_col_type(values):
    """Bounded-sample type inference over raw openpyxl cell values -- there
    is no pandas Series/dtype to inspect once streaming via
    `iter_rows(values_only=True)`, so this mirrors the REST connector's own
    `_infer_type()` shape instead."""
    non_null = [v for v in values if v is not None]
    if not non_null:
        return String()
    if all(isinstance(v, bool) for v in non_null):
        return Boolean()
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in non_null):
        if all(isinstance(v, int) and not isinstance(v, bool) for v in non_null):
            return Integer()
        return Float()
    if all(isinstance(v, (_datetime, date)) for v in non_null):
        return DateTime()
    return String()


def _approx_row_bytes(record: dict) -> int:
    """Approximate wire/storage bytes for one Excel row dict (D-16) -- there
    is no HTTP response body or DataFrame `memory_usage()` to measure here,
    so this sums a UTF-8 text length per non-null cell. Deliberately
    approximate, matching the plan's "tracking approximate bytes" wording
    for the Excel branch."""
    total = 0
    for v in record.values():
        if v is not None:
            total += len(str(v).encode("utf-8", "ignore"))
    return total


class FileConnector(DataSourceConnector):
    def __init__(self, ds):
        opts = {}
        try:
            opts = json.loads(ds.options or "{}")
        except (json.JSONDecodeError, TypeError):
            opts = {}
        self._path = opts.get("file_path")
        self._kind = opts.get("kind") or ds.source_type
        if not self._path:
            raise ConnectorError("Source file is missing — re-upload the file")
        # SECURITY: confine to the upload directory. Files are only ever placed
        # under DATA_DIR/qb_sources/<id>/ by the upload route; rejecting any
        # other path stops a crafted options.file_path from reading arbitrary
        # server files (e.g. data/.app_key, auth.json) as a queryable table.
        allowed_root = os.path.realpath(os.path.join(DATA_DIR, "qb_sources"))
        try:
            resolved = os.path.realpath(self._path)
            in_root = os.path.commonpath([allowed_root, resolved]) == allowed_root
        except (ValueError, OSError):
            in_root = False
        if not in_root:
            raise ConnectorError("file source path is outside the allowed upload directory")
        if not os.path.exists(self._path):
            raise ConnectorError("Source file is missing — re-upload the file")

    def make_engine(self):
        raise ConnectorError("File sources have no SQL engine")

    def reflect_table(self, name: str):
        raise ConnectorError("File sources do not support reflection")

    def test_connection(self) -> None:
        _pandas()
        if not os.path.exists(self._path):
            raise ConnectorError("Source file is missing")

    def list_tables(self):
        pd = _pandas()
        if self._kind == "excel":
            return list(pd.ExcelFile(self._path).sheet_names)
        # CSV: a single logical table named after the file stem.
        return [PurePath(self._path).stem]

    def fetch_table(self, name: str) -> SyncTable:
        if self._kind == "excel":
            return self._fetch_excel(name)
        return self._fetch_csv()

    def _fetch_csv(self) -> SyncTable:
        pd = _pandas()
        # (1) bounded first sample -> fresh column names/types (D-17).
        sample = pd.read_csv(self._path, nrows=CSV_SAMPLE_ROWS)
        columns = [SyncColumn(str(c), _dtype_to_sa(sample[c].dtype)) for c in sample.columns]
        st = SyncTable(columns=columns, rows=iter(()))

        def _stream():
            # QB-10/Pitfall 4: read st.columns[i].type_ LAZILY, at generator-
            # consumption time -- i.e. AFTER resolve_effective_types() has
            # already mutated it in place per its documented calling contract
            # (called BETWEEN fetch_table() and materialize_synctable()) --
            # never the freshly-inferred sample type captured above. This is
            # what lets every chunk parse against the ALREADY-RESOLVED type.
            # CR-02 (HIGH, QB-17) reviewed this path: unlike REST's
            # non-paginated single-GET (an unbounded, potentially
            # attacker/compromised-upstream-controlled remote response),
            # `pd.read_csv(..., chunksize=CSV_CHUNK_ROWS)` already bounds each
            # chunk read to CSV_CHUNK_ROWS (1000) rows before it ever reaches
            # this loop -- pandas itself never buffers past one chunk at a
            # time, so peak memory here is O(CSV_CHUNK_ROWS), not O(file
            # size). A single pathological (e.g. multi-GB) CELL within one
            # row is an accepted residual risk: the source file is uploaded
            # by an authenticated admin (not a remote/attacker-controlled
            # response), and bounding per-row would require abandoning
            # pandas' own chunked CSV reader.
            dtype_map = {c.name: _sa_type_to_pandas_dtype(c.type_) for c in st.columns}
            bytes_read = 0
            rows_yielded = 0
            skiprows = None
            while True:
                try:
                    for chunk in pd.read_csv(self._path, chunksize=CSV_CHUNK_ROWS,
                                             dtype=dtype_map, skiprows=skiprows):
                        chunk_bytes = int(chunk.memory_usage(deep=True).sum())
                        if bytes_read + chunk_bytes > MAX_SYNC_BYTES:
                            # Discard this whole chunk rather than partially
                            # yielding it (matches the REST connector's
                            # per-page discard behavior, 23-04).
                            st.truncated = True
                            return
                        bytes_read += chunk_bytes
                        chunk = chunk.where(chunk.notna(), None)
                        for row in chunk.to_dict(orient="records"):
                            yield row
                            rows_yielded += 1
                    return  # whole file consumed successfully
                except ValueError as e:
                    # CR-03 (HIGH, QB-17/QB-10 "stability"): a value past the
                    # bounded CSV_SAMPLE_ROWS sample violated the dtype
                    # pinned from that sample (e.g. an "id"/"amount" column
                    # that looked numeric in the first 500 rows but holds a
                    # stray "N/A"/text value later in the file) -- pandas
                    # raises here, INSIDE this generator, while parsing the
                    # offending chunk (before this fix that propagated
                    # uncaught and crashed the whole table sync -- a real
                    # regression vs the pre-phase whole-file pd.read_csv()
                    # with no dtype= pin, which tolerated exactly this by
                    # inferring `object`). Degrade gracefully instead: widen
                    # EVERY still-pinned column to `object` (text) -- safe,
                    # can never itself mis-parse -- and RESUME the read
                    # (skiprows past whatever has already been yielded)
                    # rather than aborting the sync. We can't reliably
                    # attribute the failure to a single column from pandas'
                    # message alone (it names the value, not the column), so
                    # this widens the whole pinned set rather than guessing.
                    if all(v == "object" for v in dtype_map.values()):
                        raise  # already fully widened -- a different, unrecoverable error
                    logger.warning(
                        "CSV dtype pin violated by data past the %s-row sample "
                        "for source file %r (%s); widening all pinned columns "
                        "to text and resuming the sync instead of aborting it",
                        CSV_SAMPLE_ROWS, self._path, e)
                    dtype_map = {k: "object" for k in dtype_map}
                    skiprows = range(1, rows_yielded + 1)

        st.rows = _stream()
        return st

    def _fetch_excel(self, name: str) -> SyncTable:
        openpyxl = _openpyxl()
        wb = openpyxl.load_workbook(self._path, read_only=True, data_only=True)
        ws = wb[name]
        row_iter = ws.iter_rows(values_only=True)
        header = list(next(row_iter, ()))
        # openpyxl expands every yielded row to the worksheet's max column.
        # A stray trailing value in a later row can therefore pad the header
        # with trailing None cells, making those values look like legitimate
        # synthetic columns and defeating the drop-and-warn contract below.
        # Preserve unnamed columns inside the real header span, but trim only
        # trailing blank cells before fixing the schema width.
        while header and header[-1] is None:
            header.pop()
        keys = [str(h) if h is not None else f"col{i}" for i, h in enumerate(header)]
        # Bounded first sample for column/type inference (D-17) -- mirrors
        # the REST connector's itertools.islice/chain peek pattern, since
        # openpyxl's read_only iterator can't be scanned in full without
        # defeating the streaming goal.
        sample_rows = list(itertools.islice(row_iter, EXCEL_SAMPLE_ROWS))
        columns = [
            SyncColumn(k, _infer_excel_col_type([r[i] if i < len(r) else None for r in sample_rows]))
            for i, k in enumerate(keys)
        ]
        st = SyncTable(columns=columns, rows=iter(()))

        def _stream():
            bytes_read = 0
            # WR-03 (LOW): the column set is fixed from the header row, but a
            # row with MORE raw values than there are header-derived `keys`
            # (e.g. a stray trailing cell past the declared header width)
            # silently drops those extra values below -- previously with no
            # log/signal anywhere. Warn once (not per-row) so it's at least
            # observable.
            warned_extra_cols = False
            try:
                for row in itertools.chain(sample_rows, row_iter):
                    if not warned_extra_cols and len(row) > len(keys):
                        logger.warning(
                            "Excel sheet %r in source file %r has a row with "
                            "%d value(s) beyond the %d header-derived "
                            "column(s) -- the extra value(s) are dropped "
                            "from every row", name, self._path,
                            len(row) - len(keys), len(keys))
                        warned_extra_cols = True
                    record = {k: (row[i] if i < len(row) else None) for i, k in enumerate(keys)}
                    row_bytes = _approx_row_bytes(record)
                    if bytes_read + row_bytes > MAX_SYNC_BYTES:
                        st.truncated = True
                        return
                    bytes_read += row_bytes
                    yield record
            finally:
                wb.close()

        st.rows = _stream()
        return st
